import io
import json
import os
import shutil
import tempfile
import unittest
from unittest.mock import MagicMock, patch
from io import BytesIO
from PIL import Image as PILImage
from docx import Document

from tool import (
    _add_answer_key_table,
    _set_default_font,
    _add_picture_fitted,
    _add_question_docx,
    _fetch_image_bytes,
    _IMG_CACHE,
    _IMG_HEADERS,
    _SESSION_COOKIES,
    _HARVESTED_LINKS,
    set_proxy,
    extract_discussion_info,
    find_discussion_link,
    build_html,
    canonical_exam_code,
    convert_to_docx,
    convert_to_html,
    is_examtopics_discussion_url,
    link_matches_question,
    load_all,
    has_good_data,
    no_link_result,
    NO_DISCUSSION_BLOCKED,
    NO_DISCUSSION_MISSING,
    normalize_exam_code,
    parse_range,
    save_progress,
    SEARCH_BLOCKED,
    SEARCH_EMPTY,
    SEARCH_ERROR,
    SEARCH_OK,
    sync_browser_session,
    upsert,
    unwrap_search_href,
    as_int,
    record_error,
)


class TestParseRange(unittest.TestCase):
    def test_range(self):
        self.assertEqual(parse_range("1-10"), (1, 10))

    def test_single(self):
        self.assertEqual(parse_range("5"), (5, 5))

    def test_invalid_text(self):
        self.assertIsNone(parse_range("abc"))

    def test_empty(self):
        self.assertIsNone(parse_range(""))

    def test_reversed(self):
        self.assertEqual(parse_range("10-1"), (10, 1))


class TestExamCodeNormalization(unittest.TestCase):
    def test_canonical_strips_separators(self):
        self.assertEqual(canonical_exam_code("SK0-005"), "sk0005")
        self.assertEqual(canonical_exam_code("FCSS_NST_SE-7.6"), "fcssnstse76")
        self.assertEqual(canonical_exam_code("H12-711_V4.0"), "h12711v40")

    def test_normalize_preserves_valid_slug_chars(self):
        self.assertEqual(normalize_exam_code("SK0-005"), "sk0-005")
        self.assertEqual(normalize_exam_code("  FCSS_NST_SE-7.6  "), "fcss_nst_se-7.6")
        self.assertEqual(normalize_exam_code("H12-711_V4.0"), "h12-711_v4.0")
        self.assertEqual(normalize_exam_code("az 104"), "az-104")


class TestLinkMatching(unittest.TestCase):
    def test_valid_discussion_slug(self):
        url = "https://www.examtopics.com/discussions/comptia/view/61482-exam-sk0-005-topic-1-question-1-discussion/"
        self.assertTrue(link_matches_question(url, "sk0-005", 1, 1))

    def test_mismatched_qnum(self):
        url = "https://www.examtopics.com/discussions/comptia/view/61482-exam-sk0-005-topic-1-question-10-discussion/"
        self.assertFalse(link_matches_question(url, "sk0-005", 1, 1))

    def test_dotted_code_matching(self):
        url = "https://www.examtopics.com/discussions/huawei/view/12345-exam-h12-711_v40-topic-1-question-2-discussion/"
        self.assertTrue(link_matches_question(url, "H12-711_V4.0", 1, 2))

    def test_non_discussion_url(self):
        self.assertFalse(link_matches_question("https://www.google.com", "sk0-005", 1, 1))
        self.assertFalse(link_matches_question(None, "sk0-005", 1, 1))

    def test_is_examtopics_discussion_url(self):
        self.assertTrue(is_examtopics_discussion_url("https://www.examtopics.com/discussions/view/1/"))
        self.assertFalse(is_examtopics_discussion_url("https://www.examtopics.com/exams/comptia/"))
        self.assertFalse(is_examtopics_discussion_url("https://example.com/discussions/"))


class TestUnwrapHref(unittest.TestCase):
    def test_unwrap_duckduckgo_uddg(self):
        raw = "https://duckduckgo.com/l/?uddg=https%3A%2F%2Fwww.examtopics.com%2Fdiscussions%2Ftest"
        urls = unwrap_search_href(raw, "https://duckduckgo.com")
        self.assertTrue(any("examtopics.com/discussions/test" in u for u in urls))

    def test_unwrap_relative_discussions(self):
        urls = unwrap_search_href("/discussions/item", "https://www.examtopics.com")
        self.assertTrue(any(u.startswith("https://www.examtopics.com/discussions/item") for u in urls))


class TestNoLinkResult(unittest.TestCase):
    def test_primary_ok_fallback_blocked_not_aborted(self):
        self.assertEqual(no_link_result(SEARCH_OK, SEARCH_BLOCKED), NO_DISCUSSION_MISSING)
        self.assertEqual(no_link_result(SEARCH_EMPTY, SEARCH_BLOCKED), NO_DISCUSSION_MISSING)

    def test_both_blocked_aborts(self):
        self.assertEqual(no_link_result(SEARCH_BLOCKED, SEARCH_BLOCKED), NO_DISCUSSION_BLOCKED)
        self.assertEqual(no_link_result(SEARCH_BLOCKED, SEARCH_ERROR), NO_DISCUSSION_BLOCKED)

    def test_primary_blocked_fallback_ok_not_aborted(self):
        self.assertEqual(no_link_result(SEARCH_BLOCKED, SEARCH_OK), NO_DISCUSSION_MISSING)


class TestLoadAllDataSafety(unittest.TestCase):
    """load_all KHONG duoc coi file hong la rong roi de save_progress ghi de
    -> mat toan bo du lieu cu. File hong phai duoc cach ly (doi ten) chu khong
    bi xoa.
    """

    def setUp(self):
        import examtopic.parser as P
        self.P = P
        self.tmp = tempfile.mkdtemp()
        self._old_dir = P.OUTPUT_DIR
        P.OUTPUT_DIR = self.tmp

    def tearDown(self):
        self.P.OUTPUT_DIR = self._old_dir
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _path(self, name="x_questions.json"):
        return os.path.join(self.tmp, name)

    def test_missing_file_returns_empty(self):
        self.assertEqual(load_all(self._path()), [])

    def test_corrupt_file_is_quarantined_not_destroyed(self):
        import glob
        path = self._path()
        good = [{"topic": 1, "question_num": i, "question": f"q{i}"} for i in range(1, 51)]
        with open(path, "w", encoding="utf-8") as f:
            json.dump(good, f)
        with open(path, encoding="utf-8") as f:
            raw = f.read()
        with open(path, "w", encoding="utf-8") as f:
            f.write(raw[: len(raw) // 2])  # cat nua -> JSON hong

        data = load_all(path)
        self.assertEqual(data, [])
        # File hong phai duoc giu lai duoi dang .corrupt-*
        quarantined = glob.glob(path + ".corrupt-*")
        self.assertTrue(quarantined, "file hong phai duoc cach ly")
        with open(quarantined[0], encoding="utf-8") as f:
            self.assertGreater(len(f.read()), 0,
                               "noi dung file hong phai con nguyen de cuu")

    def test_non_list_json_is_quarantined(self):
        import glob
        path = self._path()
        with open(path, "w", encoding="utf-8") as f:
            f.write('{"not": "a list"}')
        self.assertEqual(load_all(path), [])
        self.assertTrue(glob.glob(path + ".corrupt-*"))

    def test_save_progress_is_readable_after_write(self):
        path = self._path()
        data = [{"topic": 1, "question_num": 1, "question": "q"}]
        self.assertTrue(save_progress(data, "x_questions.json"))
        with open(path, encoding="utf-8") as f:
            self.assertEqual(json.load(f), data)


class TestUpsertAndOrder(unittest.TestCase):
    def test_upsert_maintains_sorted_order(self):
        data = []
        upsert(data, {"topic": 1, "question_num": 20, "question": "Q20"})
        upsert(data, {"topic": 1, "question_num": 5, "question": "Q5"})
        upsert(data, {"topic": 2, "question_num": 1, "question": "T2Q1"})
        upsert(data, {"topic": 1, "question_num": 1, "question": "Q1"})

        expected = [(1, 1), (1, 5), (1, 20), (2, 1)]
        actual = [(q["topic"], q["question_num"]) for q in data]
        self.assertEqual(actual, expected)

    def test_upsert_overwrites_existing(self):
        data = [{"topic": 1, "question_num": 5, "question": "Old Q5"}]
        upsert(data, {"topic": 1, "question_num": 5, "question": "New Q5"})
        self.assertEqual(len(data), 1)
        self.assertEqual(data[0]["question"], "New Q5")


class TestNegativeImageCache(unittest.TestCase):
    def test_failed_image_sets_negative_cache(self):
        url = "https://invalid-non-existent-domain-test.com/fake.png"
        _IMG_CACHE.clear()
        with patch("examtopic.config.httpx.get", side_effect=Exception("Network failure")) as mock_get:
            res = _fetch_image_bytes(url)
            self.assertIsNone(res)
            self.assertIn(url, _IMG_CACHE)
            self.assertIsNone(_IMG_CACHE[url])
            self.assertEqual(mock_get.call_count, 1)

            with patch("examtopic.config.httpx.get") as mock_get2:
                res2 = _fetch_image_bytes(url)
                self.assertIsNone(res2)
                mock_get2.assert_not_called()


class TestBuildHtml(unittest.TestCase):
    def test_no_double_escaping_in_payload(self):
        questions = [
            {
                "exam_code": "test",
                "topic": 1,
                "question_num": 1,
                "question": "What is & why?",
                "question_images": ["https://img.test.com/q.png?a=1&b=2"],
                "options": [
                    {"letter": "A", "text": "Opt & text", "images": [], "is_correct": True}
                ],
                "suggested_answers": ["A"],
                "answers": []
            }
        ]
        html_out = build_html(questions, "test", embed_images=False)
        self.assertNotIn("&amp;amp;", html_out)
        self.assertIn("https://img.test.com/q.png?a=1&b=2", html_out)

    def test_multi_topic_labels(self):
        questions = [
            {"topic": 1, "question_num": 1, "question": "T1Q1", "options": [], "suggested_answers": []},
            {"topic": 2, "question_num": 1, "question": "T2Q1", "options": [], "suggested_answers": []}
        ]
        html_out = build_html(questions, "test", embed_images=False)
        self.assertIn("Topic 1", html_out)
        self.assertIn("Topic 2", html_out)
        self.assertIn("T1 #1", html_out)
        self.assertIn("T2 #1", html_out)


class TestDocxImageFitted(unittest.TestCase):
    def test_add_picture_fitted(self):
        doc = Document()
        im = PILImage.new("RGB", (50, 50), color="blue")
        buf = BytesIO()
        im.save(buf, format="PNG")
        buf.seek(0)

        success = _add_picture_fitted(doc, buf, max_width_inches=5.5)
        self.assertTrue(success)


class TestAnswerKeyTable(unittest.TestCase):
    def test_answer_key_creates_compact_table(self):
        doc = Document()
        questions = [
            {"topic": 1, "question_num": 1, "question": "Q1", "suggested_answers": ["A"]},
            {"topic": 1, "question_num": 2, "question": "Q2", "suggested_answers": ["B", "C"]},
            {"topic": 1, "question_num": 3, "question": "Q3", "suggested_answers": ["D"]},
            {"topic": 1, "question_num": 4, "question": "Q4", "suggested_answers": []},
            {"topic": 1, "question_num": 5, "question": "Q5", "community_most_voted": ["C"]},
        ]
        _add_answer_key_table(doc, questions)
        self.assertEqual(len(doc.tables), 1)
        table = doc.tables[0]
        self.assertEqual(len(table.columns), 8)
        self.assertEqual(table.rows[0].cells[0].text, "Câu")
        self.assertEqual(table.rows[0].cells[1].text, "ĐA")


class TestBrowserSessionSync(unittest.TestCase):
    def test_sync_browser_session_updates_headers_and_cookies(self):
        mock_page = MagicMock()
        mock_page.evaluate.return_value = "Mozilla/5.0 (Custom macOS UA) TestAgent/1.0"
        mock_page.context.cookies.return_value = [
            {"name": "cf_clearance", "value": "test_cf_token"},
            {"name": "session_id", "value": "123456"}
        ]
        sync_browser_session(mock_page)
        self.assertEqual(_IMG_HEADERS["User-Agent"], "Mozilla/5.0 (Custom macOS UA) TestAgent/1.0")
        self.assertEqual(_SESSION_COOKIES.get("cf_clearance"), "test_cf_token")
        self.assertEqual(_SESSION_COOKIES.get("session_id"), "123456")


class TestCommunityMostVotedDisplay(unittest.TestCase):
    def test_html_displays_most_voted_chip_when_different(self):
        questions = [
            {
                "topic": 1,
                "question_num": 1,
                "question": "What is x?",
                "options": [],
                "suggested_answers": ["A"],
                "community_most_voted": ["B"]
            }
        ]
        html_out = build_html(questions, "test", embed_images=False)
        self.assertIn("Most Voted: B", html_out)

    def test_docx_displays_community_voted_when_different(self):
        doc = Document()
        q = {
            "topic": 1,
            "question_num": 1,
            "question": "What is x?",
            "options": [],
            "suggested_answers": ["A"],
            "community_most_voted": ["B"]
        }
        _add_question_docx(doc, q, 1)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("Cộng đồng: B", full_text)




class TestDiscussionInfoExtraction(unittest.TestCase):
    def test_extract_discussion_info_valid(self):
        url = "https://www.examtopics.com/discussions/comptia/view/71586-exam-sk0-005-topic-1-question-3-discussion/"
        info = extract_discussion_info(url)
        self.assertEqual(info, ("sk0005", 1, 3))

    def test_extract_discussion_info_dotted(self):
        url = "https://www.examtopics.com/discussions/huawei/view/12345-exam-h12-711_v40-topic-2-question-15-discussion/"
        info = extract_discussion_info(url)
        self.assertEqual(info, ("h12711v40", 2, 15))

    def test_extract_discussion_info_invalid(self):
        self.assertIsNone(extract_discussion_info("https://example.com"))
        self.assertIsNone(extract_discussion_info("https://www.examtopics.com/exams/comptia/"))


class TestHarvestedLinksCache(unittest.TestCase):
    def test_harvested_links_skips_search(self):
        _HARVESTED_LINKS.clear()
        key = ("sk0005", 1, 99)
        test_url = "https://www.examtopics.com/discussions/comptia/view/99999-exam-sk0-005-topic-1-question-99-discussion/"
        _HARVESTED_LINKS[key] = test_url

        mock_page = MagicMock()
        result = find_discussion_link(mock_page, "sk0-005", 1, 99)
        self.assertEqual(result, test_url)
        mock_page.goto.assert_not_called()


class TestFindDiscussionLinkQueryFallback(unittest.TestCase):
    """Bug that: cau 108 ccaak ton tai that nhung tool bao "khong co discussion".

    Nguyen nhan: dang query chinh luon kem `site:examtopics.com`. Search engine
    khi do tu dong SUA ma de (ccaak -> ccsk) va tra ve cau cung so cua ma de
    KHAC -> khong link nao khop -> bao thieu cau. Do tren ccaak: cau 108, 98,
    85, 57, 86 deu ton tai, chi tim thay khi BO `site:` va boc ma de trong ngoac
    kep. Test nay khoa HANH VI: dang query chinh that bai thi phai thu tiep dang
    du phong, va chi nhan link khop dung ma de + cau.
    """

    TARGET = "https://www.examtopics.com/discussions/confluent/view/382176-exam-ccaak-topic-1-question-108-discussion/"

    def _page_returning(self, hrefs_for_query, seen):
        """Trang gia + patch search_engine: ghi lai query da thu, tra ve SERP gia.

        `search_engine` duoc patch nen khong can browser; `extract_matching_link`
        van chay THAT tren cac anchor do trang gia tra ve, nho vay test kiem tra
        dung logic so khop ma de + cau.
        """
        page = MagicMock()
        page.url = "https://duckduckgo.com/"

        def fake_search(_page, query, engine="duckduckgo"):
            seen.append((engine, query))
            return SEARCH_OK

        def query_selector_all(sel):
            anchors = []
            for h in hrefs_for_query(seen[-1][1]):
                a = MagicMock()
                a.get_attribute.return_value = h
                anchors.append(a)
            return anchors

        page.query_selector_all.side_effect = query_selector_all
        patcher = patch("examtopic.crawler.search_engine", side_effect=fake_search)
        patcher.start()
        self.addCleanup(patcher.stop)
        # Tat bo giai tat dinh trong nhom test nay: chung kiem tra duong SEARCH.
        # (Duong tat dinh co test rieng o TestDeterministicResolver.)
        p2 = patch("examtopic.resolver.resolve", return_value=None)
        p2.start()
        self.addCleanup(p2.stop)
        return page

    def test_falls_back_to_unquoted_query_when_site_query_yields_other_exam(self):
        """Dang `site:` chi ra cau 108 cua ma de khac -> phai thu dang du phong."""
        other_exam = "https://www.examtopics.com/discussions/isaca/view/92443-exam-ccak-topic-1-question-108-discussion/"

        def hrefs(query):
            if 'site:examtopics.com' in query:
                return [other_exam]
            return [other_exam, self.TARGET]

        seen = []
        page = self._page_returning(hrefs, seen)
        _HARVESTED_LINKS.clear()
        result = find_discussion_link(page, "ccaak", 1, 108)
        self.assertEqual(result, self.TARGET)
        self.assertTrue(any('"ccaak"' in q for _, q in seen),
                        "dang du phong phai boc ma de trong ngoac kep")
        self.assertTrue(any('site:examtopics.com' in q for _, q in seen),
                        "dang query chinh phai duoc thu truoc")

    def test_does_not_return_link_of_different_exam(self):
        """Link cua ma de khac KHONG duoc coi la khop (ke ca khi cung so cau)."""
        other_exam = "https://www.examtopics.com/discussions/isaca/view/92443-exam-ccak-topic-1-question-108-discussion/"
        seen = []
        page = self._page_returning(lambda q: [other_exam], seen)
        _HARVESTED_LINKS.clear()
        result = find_discussion_link(page, "ccaak", 1, 108)
        self.assertIs(result, NO_DISCUSSION_MISSING)

    def test_tries_second_query_form_for_every_engine(self):
        """Moi dang query phai duoc thu o CA DuckDuckGo lan Google."""
        seen = []
        page = self._page_returning(lambda q: [], seen)
        _HARVESTED_LINKS.clear()
        find_discussion_link(page, "ccaak", 1, 108)
        engines = {e for e, _ in seen}
        self.assertEqual(engines, {"duckduckgo", "google"})
        self.assertEqual(len(seen), 4, "2 dang query x 2 engine")

    def test_reports_blocked_only_when_both_engines_blocked(self):
        page = MagicMock()
        page.url = "https://duckduckgo.com/"
        page.query_selector_all.return_value = []
        page.keyboard.type.side_effect = None
        page.query_selector.return_value = None
        _HARVESTED_LINKS.clear()
        with patch("examtopic.crawler.search_engine", return_value=SEARCH_BLOCKED):
            result = find_discussion_link(page, "ccaak", 1, 108)
        self.assertIs(result, NO_DISCUSSION_BLOCKED)


class TestDeterministicResolver(unittest.TestCase):
    """Bo giai TAT DINH: qnum -> URL discussion khong dung search engine.

    Cơ che (do tren du lieu that cua ccaak):
      - `/exams/<bat-ky>/<code>/view/1/` lo ra 10 question_id dau (category
        trong URL la cosmetic nen khong can biet vendor).
      - `/ajax/discussion/exam-question/<question_id>` tra ve discussion_id +
        title => vua lay duoc URL vua VERIFY duoc dung de/topic/cau.
      - question_id lien mach theo khoi: ccaak q1..q54 -> 949311..949364,
        q55..q109 -> 977877..977931.

    Nho do: 1 anchor moi khoi la giai duoc ca khoi, va doan sai bi loai bo nho
    doi chieu title -- khong bao gio tra ve link cua cau khac.
    """

    def setUp(self):
        from examtopic import resolver
        self.R = resolver
        self.tmp = tempfile.mkdtemp()
        self._old = resolver.OUTPUT_DIR
        resolver.OUTPUT_DIR = self.tmp
        resolver.clear_anchors()
        self.addCleanup(self._restore)

    def _restore(self):
        self.R.OUTPUT_DIR = self._old
        self.R.clear_anchors()
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_parse_title_extracts_code_topic_question(self):
        self.assertEqual(self.R.parse_title("Exam CCAAK topic 1 question 108 discussion"),
                         ("ccaak", 1, 108))
        self.assertEqual(self.R.parse_title("Exam SC-300 topic 2 question 131 discussion"),
                         ("sc300", 2, 131))
        self.assertIsNone(self.R.parse_title("Exam CCAAK topic 1 question 108"))

    def test_question_id_from_html(self):
        html = ('<div class="question-body mt-3 pt-3 border-top" '
                'data-id="977931"><p>hi</p></div>')
        self.assertEqual(self.R.question_id_from_html(html), 977931)
        self.assertIsNone(self.R.question_id_from_html("<html>nothing</html>"))
        self.assertIsNone(self.R.question_id_from_html(""))

    def test_discussion_url_uses_real_category(self):
        url = self.R.discussion_url("382176", "ccaak", 1, 108, category="confluent")
        self.assertEqual(
            url,
            "https://www.examtopics.com/discussions/confluent/view/"
            "382176-exam-ccaak-topic-1-question-108-discussion/")

    def test_discussion_url_strips_dots_from_exam_slug(self):
        """Dau CHAM trong slug ma de lam server tra 404 -> URL chet.

        Do tren trang that: `/view/<id>-exam-h12-711_v4.0-...` -> 404, con
        `/view/<id>-exam-h12711v40-...` -> 200. Truoc day ham nay dung
        `normalize_exam_code` (giu `_` va `.`) nen moi ma de co dau cham
        (H12-711_V4.0, H12-411_V2.0, FCSS_NST_SE-7.6) sinh URL chet, khien
        tool bao "khong lay duoc cau hoi" cho cau hoi ton tai that.
        """
        for code, expected_slug in (
            ("H12-711_V4.0", "h12711v40"),
            ("H12-411_V2.0", "h12411v20"),
            ("FCSS_NST_SE-7.6", "fcssnstse76"),
            ("sk0-005", "sk0005"),
        ):
            url = self.R.discussion_url("325747", code, 1, 1, category="huawei")
            tail = url.split("/view/", 1)[1]
            self.assertIn(f"-exam-{expected_slug}-topic-1-question-1-discussion/", url)
            self.assertNotIn(".", tail, f"slug ma de khong duoc chua dau cham: {url}")
            self.assertNotIn("_", tail, f"slug ma de khong duoc chua gach duoi: {url}")

    def test_discussion_url_category_slot_accepts_any_form(self):
        """Category la cosmetic (dien gi cung 200) - chuan hoa cho gon, khong hong URL."""
        url = self.R.discussion_url("325747", "H12-711_V4.0", 1, 1, category="huawei")
        self.assertTrue(url.startswith(
            "https://www.examtopics.com/discussions/huawei/view/"))

    def test_resolve_uses_nearest_anchor_and_verifies(self):
        """q18 = q1 + 17 => 949328, va title phai khop moi duoc tra ve."""
        self.R.record_anchor("ccaak", 1, 1, 949311)

        def fake_ajax(qid):
            if qid == 949328:
                return "314606", "Exam CCAAK topic 1 question 18 discussion"
            return None

        with patch.object(self.R, "ajax_lookup", side_effect=fake_ajax):
            url = self.R.resolve("ccaak", 1, 18)
        self.assertIn("314606-exam-ccaak-topic-1-question-18-discussion", url)

    def test_resolve_rejects_wrong_question(self):
        """Doan ra cau KHAC thi phai bo qua, khong tra ve link sai."""
        self.R.record_anchor("ccaak", 1, 1, 949311)

        def fake_ajax(qid):
            # Tra ve title cua cau 999 -> khong khop cau 18 dang tim.
            return "111111", "Exam CCAAK topic 1 question 999 discussion"

        with patch.object(self.R, "ajax_lookup", side_effect=fake_ajax):
            self.assertIsNone(self.R.resolve("ccaak", 1, 18))

    def test_resolve_rejects_different_exam(self):
        """Cung so cau nhung khac ma de thi khong duoc nhan."""
        self.R.record_anchor("ccaak", 1, 1, 949311)

        def fake_ajax(qid):
            return "92443", "Exam CCAK topic 1 question 18 discussion"

        with patch.object(self.R, "ajax_lookup", side_effect=fake_ajax):
            self.assertIsNone(self.R.resolve("ccaak", 1, 18))

    def test_resolve_returns_none_without_anchor(self):
        """Khong co anchor nao -> tra None de caller rot ve search."""
        with patch.object(self.R, "seed_from_exam_page", return_value=0):
            self.assertIsNone(self.R.resolve("ccaak", 1, 108))

    def test_one_anchor_unlocks_whole_block(self):
        """1 anchor o khoi 2 phai giai duoc moi cau trong khoi do.

        Khoi 2 cua ccaak: q55 -> 977877 ... q109 -> 977931 (lien mach), nen
        q78 = 977877+23 = 977900 va q85 = 977877+30 = 977907.
        """
        self.R.record_anchor("ccaak", 1, 109, 977931)
        table = {977900: ("382212", 78), 977907: ("382210", 85),
                 977931: ("382204", 109)}

        def fake_ajax(qid):
            if qid in table:
                did, q = table[qid]
                return did, f"Exam CCAAK topic 1 question {q} discussion"
            return None

        with patch.object(self.R, "ajax_lookup", side_effect=fake_ajax):
            self.assertIn("382212-exam-ccaak-topic-1-question-78-discussion",
                          self.R.resolve("ccaak", 1, 78))
            self.assertIn("382210-exam-ccaak-topic-1-question-85-discussion",
                          self.R.resolve("ccaak", 1, 85))

    def test_anchors_persist_across_runs(self):
        """Anchor phai song qua lan chay sau (doc lai tu dia)."""
        self.R.record_anchor("ccaak", 1, 108, 977930)
        self.R.clear_anchors()
        self.R.load_anchors()
        self.assertEqual(self.R.ANCHORS["ccaak"][(1, 108)], 977930)

    def test_corrupt_anchor_file_does_not_crash(self):
        with open(self.R._anchor_path(), "w", encoding="utf-8") as f:
            f.write("{not json")
        self.R.clear_anchors()
        self.assertEqual(self.R.load_anchors(), {})

    def test_seed_parses_exam_page(self):
        """Trang exam that co dang: Question #N ... Topic M ... data-id."""
        html = (
            '<div class="card-header">Question #1 '
            '<span class="question-title-topic pull-right">Topic 1</span></div>'
            '<div class="card-body question-body" data-id="949311">q</div>'
            '<div class="card-header">Question #2 '
            '<span class="question-title-topic pull-right">Topic 1</span></div>'
            '<div class="card-body question-body" data-id="949312">q</div>'
        )
        with patch.object(self.R, "_http_get", return_value=html):
            added = self.R.seed_from_exam_page("ccaak")
        self.assertEqual(added, 2)
        self.assertEqual(self.R.ANCHORS["ccaak"][(1, 1)], 949311)
        self.assertEqual(self.R.ANCHORS["ccaak"][(1, 2)], 949312)

    def test_find_discussion_link_prefers_deterministic_path(self):
        """Co link tat dinh thi KHONG duoc goi search engine nao."""
        url = ("https://www.examtopics.com/discussions/confluent/view/"
               "382176-exam-ccaak-topic-1-question-108-discussion/")
        _HARVESTED_LINKS.clear()
        page = MagicMock()
        with patch("examtopic.resolver.resolve", return_value=url), \
             patch("examtopic.crawler.search_engine") as mock_search:
            result = find_discussion_link(page, "ccaak", 1, 108)
        self.assertEqual(result, url)
        mock_search.assert_not_called()


class TestProxyConfiguration(unittest.TestCase):
    def test_set_proxy(self):
        set_proxy("http://127.0.0.1:8888")
        from examtopic.config import _PROXY as conf_proxy
        self.assertEqual(conf_proxy, "http://127.0.0.1:8888")
        set_proxy(None)

    def test_http_fallback_sees_proxy_set_after_import(self):
        """Regression: crawler tung `from .config import _PROXY` (copy gia tri
        luc import) nen set_proxy() sau do khong den duoc HTTP fallback -> am
        tham bo qua proxy, lo IP that.

        Test nay kiem tra HANH VI quan sat duoc: gia tri `proxy=` thuc su truyen
        vao httpx.get() cua load_discussion_via_http, chu khong chi doc thuoc
        tinh noi bo.
        """
        from examtopic.crawler import load_discussion_via_http
        set_proxy("http://127.0.0.1:7777")
        try:
            tab = MagicMock()
            tab.query_selector.return_value = MagicMock()
            resp = MagicMock(status_code=200, text="<div class='discussion-header-container'></div>")
            with patch("examtopic.crawler.httpx.get", return_value=resp) as mock_get:
                load_discussion_via_http(tab, "https://x.test")
            self.assertEqual(mock_get.call_count, 1)
            self.assertEqual(
                mock_get.call_args.kwargs.get("proxy"), "http://127.0.0.1:7777",
                "HTTP fallback phai truyen proxy da set vao httpx.get",
            )
        finally:
            set_proxy(None)

    def test_http_fallback_has_no_proxy_by_default(self):
        """Mac dinh khong set proxy thi httpx.get phai nhan proxy=None."""
        from examtopic.crawler import load_discussion_via_http
        set_proxy(None)
        tab = MagicMock()
        tab.query_selector.return_value = MagicMock()
        resp = MagicMock(status_code=200, text="<div class='discussion-header-container'></div>")
        with patch("examtopic.crawler.httpx.get", return_value=resp) as mock_get:
            load_discussion_via_http(tab, "https://x.test")
        self.assertIsNone(mock_get.call_args.kwargs.get("proxy"))


class TestSearchFastPath(unittest.TestCase):
    """Search co 2 duong: URL truc tiep (nhanh) va go tay (du phong chong CAPTCHA).

    Do tren trang that: go tay 12.9s/lan, URL truc tiep 1.8-3.7s va 22/22 lan
    lien tiep khong bi chan. Nhung duong nhanh KHONG duoc phep lam mat kha nang
    chong CAPTCHA cua ban cu -> phai tu dong rot ve go tay khi bi chan.
    """

    def _cfg(self):
        from examtopic.config import SEARCH_ENGINES
        return SEARCH_ENGINES

    def test_every_engine_has_url_template(self):
        """Thieu template thi duong nhanh tat am tham -> phai co test khoa lai."""
        for name, cfg in self._cfg().items():
            self.assertIn("url_template", cfg, f"{name} thieu url_template")
            self.assertIn("{query}", cfg["url_template"],
                          f"{name} template phai co cho noi query")

    def test_url_template_is_valid_and_encodes_query(self):
        from urllib.parse import urlparse, quote_plus
        for name, cfg in self._cfg().items():
            url = cfg["url_template"].format(query=quote_plus("exam pl-300 topic 1"))
            parsed = urlparse(url)
            self.assertEqual(parsed.scheme, "https", f"{name} phai la https")
            self.assertTrue(parsed.netloc, f"{name} phai co host")
            self.assertIn("pl-300", url)

    def test_fast_path_used_when_it_works(self):
        """Duong nhanh chay duoc -> KHONG duoc go tay (tiet kiem ~10s/lan)."""
        from examtopic import crawler
        page = MagicMock()
        page.url = "https://duckduckgo.com/?q=x"
        page.query_selector.return_value = None
        with patch.object(crawler, "_search_via_url", return_value=SEARCH_OK) as fast, \
             patch.object(crawler, "_search_by_typing") as typing:
            result = crawler.search_engine(page, "exam x topic 1 question 1", "duckduckgo")
        self.assertEqual(result, SEARCH_OK)
        fast.assert_called_once()
        typing.assert_not_called()

    def test_falls_back_to_typing_when_url_path_blocked(self):
        """URL truc tiep bi chan -> PHAI rot ve go tay, khong duoc bo cuoc."""
        from examtopic import crawler
        page = MagicMock()
        with patch.object(crawler, "_search_via_url", return_value=None), \
             patch.object(crawler, "_search_by_typing", return_value=SEARCH_OK) as typing:
            result = crawler.search_engine(page, "exam x topic 1 question 1", "duckduckgo")
        self.assertEqual(result, SEARCH_OK)
        typing.assert_called_once()

    def test_blocked_result_is_not_treated_as_failure_of_fast_path(self):
        """`SEARCH_BLOCKED` la ket qua THAT, khong phai 'duong nhanh hong'.

        Neu coi no la None thi moi lan bi chan se ton them 13s go tay vo ich.
        """
        from examtopic import crawler
        page = MagicMock()
        with patch.object(crawler, "_search_via_url", return_value=SEARCH_BLOCKED), \
             patch.object(crawler, "_search_by_typing") as typing:
            result = crawler.search_engine(page, "q", "duckduckgo")
        self.assertEqual(result, SEARCH_BLOCKED)
        typing.assert_not_called()

    def test_unknown_engine_still_errors(self):
        from examtopic import crawler
        self.assertEqual(crawler.search_engine(MagicMock(), "q", "khong-ton-tai"), SEARCH_ERROR)

    def test_direct_goto_does_not_wait_load_twice(self):
        """`goto` da cho domcontentloaded thi khong duoc wait them lan nua.

        Do duoc: cho them lan nua ton ~4s moi lan search.
        """
        from examtopic import crawler
        page = MagicMock()
        page.query_selector.return_value = None
        with patch.object(crawler, "_read_search_result") as read, \
             patch.object(crawler, "_search_is_blocked", return_value=False):
            read.return_value = SEARCH_OK
            crawler._search_via_url(page, self._cfg()["duckduckgo"], "q", "duckduckgo")
        self.assertEqual(page.goto.call_args.kwargs.get("wait_until"), "domcontentloaded")
        self.assertFalse(read.call_args.kwargs.get("wait_load", True),
                         "khong duoc wait_for_load_state lan nua")

    def test_typing_path_still_waits_for_load(self):
        """Duong go tay thi van phai cho load (goto o day dung wait_until mac dinh)."""
        from examtopic import crawler
        page = MagicMock()
        # Phai co o tim kiem that, neu khong ham thoat som truoc khi toi buoc cho load.
        page.query_selector.return_value = MagicMock()
        with patch.object(crawler, "safe_goto", return_value=True), \
             patch.object(crawler, "_search_is_blocked", return_value=False), \
             patch.object(crawler, "_read_search_result", return_value=SEARCH_OK) as read:
            crawler._search_by_typing(page, self._cfg()["duckduckgo"], "q", "duckduckgo")
        read.assert_called_once()
        self.assertTrue(read.call_args.kwargs.get("wait_load", True))


class TestLoadFullComments(unittest.TestCase):
    """Trang discussion chi render san ~20-25 binh luan dau; phan con lai nam
    sau nut "Load full discussion..." va chi lay duoc qua AJAX load-complete.

    Truoc ban va nay crawler khong bao gio goi buoc nay -> cac cau soi noi bi
    cat cut. Do tren du lieu that: 10/17 cau co >=20 binh luan bi mat trung
    binh ~30-50% noi dung (co cau hien 20/78 binh luan).

    Cac test duoi day khoa HANH VI: co goi AJAX khi co nut, KHONG goi khi
    khong co nut, va khong lam hong du lieu khac khi AJAX loi.
    """

    def _tab(self, has_button=True, did="56676", comments=20):
        """Tab gia: co/khong nut load-full-discussion, tra ve discussion-id."""
        tab = MagicMock()

        def query_selector(sel):
            if sel == '.load-full-discussion-button':
                return MagicMock() if has_button else None
            return MagicMock()

        tab.query_selector.side_effect = query_selector
        # evaluate() dau tien la _extract_discussion_id -> tra data-discussion-id.
        # evaluate() sau do la buoc bom fragment -> tra True (bom thanh cong).
        tab.evaluate.side_effect = [did, True]
        # _count_comments truoc/sau khi bom.
        tab.eval_on_selector_all.return_value = comments
        return tab

    def test_skips_ajax_when_no_load_button(self):
        """Khong co nut => trang da day du, khong ton request nao."""
        from examtopic.crawler import load_full_comments
        tab = self._tab(has_button=False)
        with patch("examtopic.crawler.httpx.get") as mock_get:
            result = load_full_comments(tab, "https://x.test/view/56676-exam-a-topic-1-question-1-discussion/")
        self.assertIsNone(result)
        mock_get.assert_not_called()

    def test_calls_ajax_when_load_button_present(self):
        """Co nut => phai goi AJAX load-complete voi dung discussion-id."""
        from examtopic.crawler import load_full_comments
        tab = self._tab(has_button=True, did="56676")
        resp = MagicMock(status_code=200, text=(
            '<div class="outer-discussion-container">'
            '<div class="comment-content">a</div></div>'
        ))
        with patch("examtopic.crawler.httpx.get", return_value=resp) as mock_get:
            load_full_comments(tab, "https://x.test/view/56676-exam-a-topic-1-question-1-discussion/")
        self.assertEqual(mock_get.call_count, 1)
        self.assertEqual(mock_get.call_args.kwargs.get("params"), {"discussion-id": "56676"})

    def test_ajax_failure_keeps_existing_dom(self):
        """AJAX loi => tra None va KHONG bom gi vao DOM (khong mat du lieu cu)."""
        from examtopic.crawler import load_full_comments
        tab = self._tab(has_button=True)
        resp = MagicMock(status_code=500, text="")
        with patch("examtopic.crawler.httpx.get", return_value=resp):
            result = load_full_comments(tab, "https://x.test/view/56676-exam-a-topic-1-question-1-discussion/")
        self.assertIsNone(result)
        # evaluate chi duoc goi 1 lan (doc discussion-id), khong bom fragment.
        self.assertEqual(tab.evaluate.call_count, 1)

    def test_fragment_without_container_is_rejected(self):
        """Fragment khong co .outer-discussion-container => coi nhu that bai."""
        from examtopic.crawler import load_full_comments
        tab = self._tab(has_button=True)
        tab.evaluate.side_effect = ["56676", False]
        resp = MagicMock(status_code=200, text="<div>khong co container</div>")
        with patch("examtopic.crawler.httpx.get", return_value=resp):
            result = load_full_comments(tab, "https://x.test/view/56676-exam-a-topic-1-question-1-discussion/")
        self.assertIsNone(result)

    def test_derives_discussion_id_from_url_when_dom_lacks_it(self):
        """DOM khong co data-discussion-id => rot ve id trong slug URL."""
        from examtopic.crawler import load_full_comments
        tab = MagicMock()
        tab.query_selector.return_value = MagicMock()
        tab.evaluate.side_effect = [None, True]
        tab.eval_on_selector_all.return_value = 5
        resp = MagicMock(status_code=200, text='<div class="outer-discussion-container"></div>')
        with patch("examtopic.crawler.httpx.get", return_value=resp) as mock_get:
            load_full_comments(tab, "https://www.examtopics.com/discussions/ms/view/56676-exam-ms-700-topic-1-question-21-discussion/")
        self.assertEqual(mock_get.call_args.kwargs.get("params"), {"discussion-id": "56676"})

    def test_uses_proxy_set_after_import(self):
        """Proxy set sau khi import van phai ap dung (cung loi nhu HTTP fallback)."""
        from examtopic.crawler import load_full_comments
        set_proxy("http://127.0.0.1:6666")
        try:
            tab = self._tab(has_button=True)
            resp = MagicMock(status_code=200, text='<div class="outer-discussion-container"></div>')
            with patch("examtopic.crawler.httpx.get", return_value=resp) as mock_get:
                load_full_comments(tab, "https://x.test/view/56676-exam-a-topic-1-question-1-discussion/")
            self.assertEqual(mock_get.call_args.kwargs.get("proxy"), "http://127.0.0.1:6666")
        finally:
            set_proxy(None)

    def test_returns_before_and_after_counts(self):
        """Tra ve (truoc, sau) de log chung minh da nap them binh luan."""
        from examtopic.crawler import load_full_comments
        tab = MagicMock()
        tab.query_selector.return_value = MagicMock()
        tab.evaluate.side_effect = ["56676", True]
        tab.eval_on_selector_all.side_effect = [20, 70]
        resp = MagicMock(status_code=200, text='<div class="outer-discussion-container"></div>')
        with patch("examtopic.crawler.httpx.get", return_value=resp):
            result = load_full_comments(tab, "https://x.test/view/56676-exam-a-topic-1-question-1-discussion/")
        self.assertEqual(result, (20, 70))

    def test_strips_scripts_from_fragment(self):
        """Fragment phai duoc loc script truoc khi bom vao DOM."""
        from examtopic.crawler import load_full_comments
        tab = MagicMock()
        tab.query_selector.return_value = MagicMock()
        tab.evaluate.side_effect = ["56676", True]
        tab.eval_on_selector_all.return_value = 5
        resp = MagicMock(status_code=200, text=(
            '<div class="outer-discussion-container"></div>'
            '<script>alert(1)</script>'
        ))
        with patch("examtopic.crawler.httpx.get", return_value=resp):
            load_full_comments(tab, "https://x.test/view/56676-exam-a-topic-1-question-1-discussion/")
        bommed = tab.evaluate.call_args.args[1]
        self.assertNotIn("<script>alert(1)</script>", bommed)
        self.assertIn("outer-discussion-container", bommed)


class TestScriptTagPreservation(unittest.TestCase):
    """_RE_SCRIPT phai GIU lai <script type="application/json"> (nguon du lieu
    community_most_voted) va STRIP cac script khac. Bien the co khoang trang
    quanh dau '=' tung bi xoa am tham -> mat du lieu vote ma khong bao loi.
    """

    def setUp(self):
        from examtopic.config import _RE_SCRIPT
        self.strip = _RE_SCRIPT.sub

    def _kept(self, html):
        return bool(self.strip('', html).strip())

    def test_keeps_canonical_json_script(self):
        self.assertTrue(self._kept('<script type="application/json">[1]</script>'))

    def test_keeps_json_script_with_space_around_equals(self):
        for html in (
            '<script type = "application/json">[1]</script>',
            '<script type= "application/json">[1]</script>',
            '<script type ="application/json" id="x">[1]</script>',
            "<script type = 'application/json'>[1]</script>",
        ):
            self.assertTrue(self._kept(html), f"bi strip nham: {html}")

    def test_keeps_real_examtopics_tally_markup(self):
        html = ('<div class="voted-answers-tally d-none">'
                '<script id="806123" type="application/json">'
                '[{"voted_answers": "D", "vote_count": 12, "is_most_voted": true}]'
                '</script></div>')
        out = self.strip('', html)
        self.assertIn("voted_answers", out)
        self.assertIn("is_most_voted", out)

    def test_strips_ordinary_scripts(self):
        for html in (
            '<script type="text/javascript">var x=1;</script>',
            '<script>var y=2;</script>',
            '<script src="https://ads.example/a.js"></script>',
        ):
            self.assertFalse(self._kept(html), f"khong strip: {html}")


class TestAsInt(unittest.TestCase):
    def test_int_passthrough(self):
        self.assertEqual(as_int(5), 5)
        self.assertEqual(as_int("7"), 7)

    def test_invalid_returns_default(self):
        self.assertEqual(as_int(None), 0)
        self.assertEqual(as_int("abc"), 0)
        self.assertEqual(as_int("abc", default=1), 1)
        self.assertEqual(as_int("", default=9), 9)


class TestUpsertRobustKeys(unittest.TestCase):
    def test_upsert_does_not_crash_on_garbage_topic(self):
        data = [{"topic": "abc", "question_num": 1, "question": "x"}]
        upsert(data, {"topic": 1, "question_num": 2, "question": "y"})
        self.assertEqual(len(data), 2)

    def test_upsert_matches_str_and_int_keys(self):
        data = [{"topic": "1", "question_num": "5", "question": "old"}]
        upsert(data, {"topic": 1, "question_num": 5, "question": "new"})
        self.assertEqual(len(data), 1)
        self.assertEqual(data[0]["question"], "new")


class TestHttpRetryBackoff(unittest.TestCase):
    def test_retries_on_429_then_succeeds(self):
        from examtopic.crawler import load_discussion_via_http
        responses = [
            MagicMock(status_code=429, text="rate limited"),
            MagicMock(status_code=200, text="<div class='discussion-header-container'></div>"),
        ]
        tab = MagicMock()
        tab.query_selector.return_value = MagicMock()
        with patch("examtopic.crawler.httpx.get", side_effect=responses) as mock_get, \
             patch("examtopic.crawler.time.sleep") as mock_sleep:
            ok = load_discussion_via_http(tab, "https://x.test")
        self.assertTrue(ok)
        self.assertEqual(mock_get.call_count, 2)
        mock_sleep.assert_called_once_with(1.0)

    def test_gives_up_after_exhausting_retries(self):
        from examtopic.crawler import load_discussion_via_http
        responses = [MagicMock(status_code=503, text="")] * 3
        tab = MagicMock()
        with patch("examtopic.crawler.httpx.get", side_effect=responses), \
             patch("examtopic.crawler.time.sleep"):
            ok = load_discussion_via_http(tab, "https://x.test")
        self.assertFalse(ok)

    def test_non_retryable_status_fails_immediately(self):
        from examtopic.crawler import load_discussion_via_http
        tab = MagicMock()
        with patch("examtopic.crawler.httpx.get",
                   return_value=MagicMock(status_code=404, text="nope")) as mock_get, \
             patch("examtopic.crawler.time.sleep") as mock_sleep:
            ok = load_discussion_via_http(tab, "https://x.test")
        self.assertFalse(ok)
        self.assertEqual(mock_get.call_count, 1)
        mock_sleep.assert_not_called()


class TestHasGoodData(unittest.TestCase):
    def test_detects_good_record(self):
        data = [{"topic": 1, "question_num": 5, "question": "real"}]
        self.assertTrue(has_good_data(data, 1, 5))

    def test_ignores_record_without_question(self):
        data = [{"topic": 1, "question_num": 5, "question": ""}]
        self.assertFalse(has_good_data(data, 1, 5))

    def test_matches_string_keys(self):
        data = [{"topic": "2", "question_num": "7", "question": "real"}]
        self.assertTrue(has_good_data(data, 2, 7))

    def test_absent_question(self):
        data = [{"topic": 1, "question_num": 5, "question": "real"}]
        self.assertFalse(has_good_data(data, 1, 6))


class TestRecordError(unittest.TestCase):
    """Ban ghi loi phai duoc tach rieng, khong tron vao file du lieu chinh."""

    def test_appends_to_error_list(self):
        errors = []
        record_error({"exam_code": "x", "topic": 1, "question_num": 2, "error": "boom"}, errors)
        record_error({"exam_code": "x", "topic": 1, "question_num": 9, "error": "nope"}, errors)
        self.assertEqual(len(errors), 2)
        self.assertEqual(errors[0]["question_num"], 2)
        self.assertNotIn("question", errors[0])


class TestSameQuestionHelper(unittest.TestCase):
    """`_same_question` quyet dinh xoa ban ghi loi cu nao khi cau da lay duoc.

    Neu so sanh sai (vd so sanh chuoi voi int), ban ghi loi cu se bi giu lai mai
    -> file .errors.json bao cau da lay duoc la loi.
    """

    def test_matches_int_keys(self):
        from tool import _same_question
        self.assertTrue(_same_question({"topic": 1, "question_num": 108}, 1, 108))

    def test_matches_string_keys(self):
        from tool import _same_question
        self.assertTrue(_same_question({"topic": "1", "question_num": "108"}, 1, 108))

    def test_rejects_other_question_or_topic(self):
        from tool import _same_question
        self.assertFalse(_same_question({"topic": 1, "question_num": 107}, 1, 108))
        self.assertFalse(_same_question({"topic": 2, "question_num": 108}, 1, 108))

    def test_rejects_non_dict(self):
        from tool import _same_question
        self.assertFalse(_same_question("boom", 1, 108))
        self.assertFalse(_same_question(None, 1, 108))

    def test_removes_only_recovered_question(self):
        """Mo phong buoc don: cau 108 lay duoc thi chi xoa ban ghi loi cua 108."""
        from tool import _same_question
        error_records = [
            {"topic": 1, "question_num": 78, "error": "old"},
            {"topic": 1, "question_num": 108, "error": "old"},
            {"topic": 1, "question_num": 98, "error": "old"},
        ]
        kept = [r for r in error_records if not _same_question(r, 1, 108)]
        self.assertEqual([r["question_num"] for r in kept], [78, 98])


class TestErrorFileResumeMerge(unittest.TestCase):
    """Lan chay lai phai GIU lai ban ghi loi cu thay vi ghi de .errors.json.

    Truoc day tool.py khoi tao `error_records = []` moi lan chay -> file
    .errors.json bi ghi de, mat am tham danh sach cau loi cua lan truoc. Nay
    phai nap lai va merge (upsert) theo (topic, question_num).
    """

    def setUp(self):
        import examtopic.parser as P
        self.P = P
        self.tmp = tempfile.mkdtemp()
        self._old_dir = P.OUTPUT_DIR
        P.OUTPUT_DIR = self.tmp

    def tearDown(self):
        self.P.OUTPUT_DIR = self._old_dir
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _path(self):
        return os.path.join(self.tmp, "x_errors.json")

    def test_second_run_keeps_first_run_errors(self):
        # Lan 1: cau 1 loi.
        recs = load_all(self._path())
        upsert(recs, {"exam_code": "x", "topic": 1, "question_num": 1, "error": "boom"})
        save_progress(recs, "x_errors.json")

        # Lan 2: nap lai roi them cau 2 loi.
        recs = load_all(self._path())
        self.assertEqual(len(recs), 1, "phai nap lai duoc ban ghi loi cu")
        upsert(recs, {"exam_code": "x", "topic": 1, "question_num": 2, "error": "nope"})
        save_progress(recs, "x_errors.json")

        with open(self._path(), encoding="utf-8") as f:
            final = json.load(f)
        self.assertEqual(len(final), 2, "khong duoc mat ban ghi loi cu")
        self.assertEqual([r["question_num"] for r in final], [1, 2])

    def test_recrawled_failure_updates_in_place(self):
        recs = [{"exam_code": "x", "topic": 1, "question_num": 1, "error": "old"}]
        save_progress(recs, "x_errors.json")

        recs = load_all(self._path())
        upsert(recs, {"exam_code": "x", "topic": 1, "question_num": 1, "error": "new"})
        save_progress(recs, "x_errors.json")

        with open(self._path(), encoding="utf-8") as f:
            final = json.load(f)
        self.assertEqual(len(final), 1, "cung cau thi cap nhat tai cho, khong nhan doi")
        self.assertEqual(final[0]["error"], "new")


class TestConvertOnlyExitCode(unittest.TestCase):
    """--convert-only phai tra ma loi khac 0 khi that bai, de script/CI khong
    tuong nham la thanh cong. Truoc day luon return ma 0.

    Chay trong tien trinh con de bat SystemExit ma khong lam chet test runner.
    """

    def _run(self, args):
        import subprocess
        import sys as _sys
        return subprocess.run(
            [_sys.executable, "tool.py", *args],
            cwd=os.path.dirname(os.path.abspath(__file__)),
            capture_output=True,
            text=True,
        )

    def test_missing_file_exits_nonzero(self):
        proc = self._run(["--convert-only", "/nonexistent/definitely-missing.json"])
        self.assertNotEqual(proc.returncode, 0, "file thieu phai tra ma loi")

    def test_no_valid_questions_exits_nonzero(self):
        tmp = tempfile.mkdtemp()
        try:
            path = os.path.join(tmp, "x_questions.json")
            with open(path, "w", encoding="utf-8") as f:
                f.write('[{"exam_code": "X", "topic": 1, "question_num": 1, "question": ""}]')
            proc = self._run(["--convert-only", path])
            self.assertNotEqual(proc.returncode, 0, "convert that bai phai tra ma loi")
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_valid_file_exits_zero(self):
        tmp = tempfile.mkdtemp()
        try:
            path = os.path.join(tmp, "x_questions.json")
            with open(path, "w", encoding="utf-8") as f:
                json.dump([{
                    "exam_code": "X", "topic": 1, "question_num": 1, "question": "q",
                    "options": [], "suggested_answers": [], "answers": [],
                }], f)
            proc = self._run(["--convert-only", path])
            self.assertEqual(proc.returncode, 0, f"convert thanh cong phai la 0: {proc.stderr}")
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


class TestConvertMalformedData(unittest.TestCase):
    """Converter khong duoc crash khi du lieu di dang (question_num/topic khong
    phai so) -- truoc day dung int() tran gay ValueError lam hong ca buoc convert.
    """

    def _write(self, data):
        tmp = tempfile.mkdtemp()
        path = os.path.join(tmp, "x_questions.json")
        with open(path, "w", encoding="utf-8") as f:
            json.dump(data, f)
        return path

    def test_non_numeric_question_num_does_not_crash(self):
        path = self._write([{
            "exam_code": "X", "topic": "1", "question_num": "abc",
            "question": "q", "options": [], "suggested_answers": [], "answers": [],
        }])
        convert_to_html(path)
        convert_to_docx(path)

    def test_none_topic_and_num_do_not_crash(self):
        path = self._write([{
            "exam_code": "X", "topic": None, "question_num": None,
            "question": "q", "options": [], "suggested_answers": [], "answers": [],
        }])
        convert_to_html(path)
        convert_to_docx(path)


class TestTopComments(unittest.TestCase):
    """Cau HOTSPOT/DRAG DROP khong co dap an A/B/C/D -> hien binh luan noi bat.

    Tin hieu dung la so luot upvote THAT cua ExamTopics (`.upvote-count`) va badge
    "Highly Voted" -- KHONG phai heuristic do tool tu nghi ra. Tool khong tu suy ra
    dap an: binh luan cong dong hay mau thuan nhau (do tren q8 that: binh luan 24
    vote chi noi "Seems correct", con binh luan 21 vote moi co dap an), nen doan
    sai con te hon de trong.
    """

    def _special(self, top=None):
        q = {"topic": 1, "question_num": 5, "question": "HOTSPOT - x",
             "options": [], "suggested_answers": [], "answers": []}
        if top is not None:
            q["top_comments"] = top
        return q

    def _normal(self, top=None):
        q = {"topic": 1, "question_num": 1, "question": "MC",
             "options": [{"letter": "A", "text": "a", "images": [], "is_correct": True}],
             "suggested_answers": ["A"], "answers": []}
        if top is not None:
            q["top_comments"] = top
        return q

    TOP = [
        {"text": "Yes/Yes/No - da kiem chung", "votes": 61, "user": "Val_0", "badge": "Highly Voted"},
        {"text": "y kien khac", "votes": 15, "user": "bob", "badge": ""},
    ]

    def test_cleaner_sorts_by_votes_and_dedupes(self):
        from examtopic.crawler import _clean_top_comments
        got = _clean_top_comments([
            {"text": "thap", "votes": 1, "user": "a", "badge": ""},
            {"text": "cao", "votes": 50, "user": "b", "badge": "Highly Voted"},
            {"text": "cao", "votes": 50, "user": "b", "badge": ""},   # trung
        ])
        self.assertEqual([c["votes"] for c in got], [50, 1])
        self.assertEqual(got[0]["text"], "cao")

    def test_cleaner_survives_junk(self):
        from examtopic.crawler import _clean_top_comments
        got = _clean_top_comments([
            None, "rac", {"text": ""}, {"text": "ok", "votes": "khong-phai-so"},
            {"text": "am", "votes": -5},
        ])
        self.assertEqual([c["text"] for c in got], ["ok", "am"])
        self.assertEqual(got[1]["votes"], 0, "vote am phai bi kep ve 0")

    def test_cleaner_handles_non_list(self):
        from examtopic.crawler import _clean_top_comments
        for bad in (None, "x", 5, {"a": 1}):
            self.assertEqual(_clean_top_comments(bad), [])

    def test_noise_filter_drops_exam_date_bragging(self):
        """Binh luan chi bao "da gap trong ky thi" phai bi loc bo.

        Do tren du lieu that ms-700 (20 cau dac biet): 15/86 (17%) binh luan noi
        bat chi la "On exam March 2023" / "On the test Nov, 12 2021"... nhung co
        vote cao vi nguoi thi roi upvote nhau, trong khi dap an that nam o binh
        luan IT vote hon. Hien chung se lam loang thong tin huu ich.
        """
        from examtopic.crawler import _is_noise_comment
        noise = [
            "On exam March 2023",
            "On the test Nov, 12 2021",
            "Was on the exam 1/3/23",
            "on exam 26-Nov-2022",
            "On test 28.04.2023 (I'm not a bot you can trust me :D)",
            "On Exam Feb 2023",
            "Seen on the exam last week",
            "Took the exam yesterday",
            "In the exam it was different",
        ]
        for t in noise:
            self.assertTrue(_is_noise_comment(t), f"phai coi la rac: {t!r}")

    def test_noise_filter_keeps_useful_comments(self):
        """Khong duoc loc nham binh luan co y kien ve dap an."""
        from examtopic.crawler import _is_noise_comment
        useful = [
            "Correct. No, No, Yes. User 3 will get the renewal notification",
            "Should be -AllowGiphy $false? https://support.microsoft.com/...",
            'Answer should be "Set-Team -AllowGiphy"',
            "For a non native speaker, animated images is equal gif, so why not allowgiphy?",
            "Correct Pureview need E5 compliance license https://example.com",
            "I passed with 900 but this question is wrong",
        ]
        for t in useful:
            self.assertFalse(_is_noise_comment(t), f"khong duoc coi la rac: {t!r}")

    def test_cleaner_applies_noise_filter(self):
        from examtopic.crawler import _clean_top_comments
        got = _clean_top_comments([
            {"text": "On exam March 2023", "votes": 99, "user": "a", "badge": "Highly Voted"},
            {"text": "Answer should be B", "votes": 3, "user": "b", "badge": ""},
        ])
        self.assertEqual([c["text"] for c in got], ["Answer should be B"],
                         "binh luan rac vote cao phai bi loai, giu lai binh luan that")

    def test_cleaner_caps_at_five(self):
        from examtopic.crawler import _clean_top_comments
        got = _clean_top_comments([{"text": f"c{i}", "votes": i} for i in range(20)])
        self.assertEqual(len(got), 5)
        self.assertEqual(got[0]["votes"], 19)

    def test_html_shows_block_for_special_question(self):
        html_out = build_html([self._special(self.TOP)], "sc-300", embed_images=False)
        self.assertIn('class="top-comments"', html_out)
        self.assertIn("61 upvote", html_out)
        self.assertIn("Highly Voted", html_out)
        self.assertIn("Val_0", html_out)
        self.assertIn("không phải đáp án chính thức", html_out)

    def test_html_hides_block_for_normal_question(self):
        """Cau co dap an A/B/C/D thi KHONG duoc hien khoi binh luan noi bat."""
        html_out = build_html([self._normal(self.TOP)], "x", embed_images=False)
        self.assertNotIn('class="top-comments"', html_out)

    def test_html_hides_block_when_no_top_comments(self):
        html_out = build_html([self._special()], "x", embed_images=False)
        self.assertNotIn('class="top-comments"', html_out)

    def test_docx_shows_block_for_special_question(self):
        doc = Document()
        _add_question_docx(doc, self._special(self.TOP), 5)
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("Bình luận nổi bật", text)
        self.assertIn("61 upvotes", text)
        self.assertIn("Highly Voted", text)
        self.assertIn("không phải đáp án chính thức", text)

    def test_docx_hides_block_for_normal_question(self):
        doc = Document()
        _add_question_docx(doc, self._normal(self.TOP), 1)
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertNotIn("Bình luận nổi bật", text)

    def test_top_comments_do_not_become_suggested_answers(self):
        """Khoa dieu quan trong nhat: tool KHONG duoc tu bien binh luan thanh dap an."""
        q = self._special(self.TOP)
        html_out = build_html([q], "x", embed_images=False)
        self.assertEqual(q["suggested_answers"], [])
        self.assertIn("Unavailable", html_out, "dap an van phai la Unavailable")


class TestAbsolutizeImages(unittest.TestCase):
    """Anh phai duoc tuyet doi hoa truoc khi luu.

    Duong nap nhanh (HTTP-first) bom HTML vao `about:blank` bang document.write,
    nen trang KHONG co base URL: `im.src` tra ve nguyen chuoi tuong doi
    `/assets/media/exam-media/...` nam trong HTML. Lop sau (`_safe_url`) coi do
    la scheme khong an toan roi loai bo -> ANH MAT AM THAM, khong co loi nao.
    Do tren du lieu that: 93 anh dang `/assets/...` trong ms700 + sk0005.
    """

    HREF = ("https://www.examtopics.com/discussions/comptia/view/"
            "71586-exam-sk0-005-topic-1-question-3-discussion/")

    def _f(self):
        from examtopic.crawler import _absolutize_images
        return _absolutize_images

    def test_relative_path_becomes_absolute(self):
        got = self._f()(["/assets/media/exam-media/04231/0000300001.png"], self.HREF)
        self.assertEqual(
            got, ["https://www.examtopics.com/assets/media/exam-media/04231/0000300001.png"])

    def test_absolute_and_data_urls_are_untouched(self):
        urls = ["https://img.examtopics.com/x.png",
                "data:image/png;base64,AAA",
                "//cdn.examtopics.com/x.png"]
        got = self._f()(urls, self.HREF)
        self.assertEqual(got[0], "https://img.examtopics.com/x.png")
        self.assertEqual(got[1], "data:image/png;base64,AAA")
        self.assertEqual(got[2], "https://cdn.examtopics.com/x.png")

    def test_idempotent(self):
        """Lan hai khong duoc doi gi (trang nap bang dieu huong that da tuyet doi)."""
        once = self._f()(["/assets/a.png"], self.HREF)
        self.assertEqual(self._f()(once, self.HREF), once)

    def test_drops_junk_and_preserves_order(self):
        got = self._f()(["", None, "/a.png", 123, "/b.png"], self.HREF)
        self.assertEqual(got, ["https://www.examtopics.com/a.png",
                               "https://www.examtopics.com/b.png"])

    def test_no_base_url_leaves_value_alone(self):
        self.assertEqual(self._f()(["/a.png"], None), ["/a.png"])
        self.assertEqual(self._f()(None, self.HREF), [])

    def test_result_passes_exporter_url_filter(self):
        """Khoa hanh vi: sau khi tuyet doi hoa, `_safe_url` cua exporter phai giu lai.

        Day moi la dieu lam anh hien ra thay vi bi loai bo am tham.
        """
        from examtopic.exporters.html import _safe_url
        raw = "/assets/media/exam-media/04231/0000300001.png"
        self.assertEqual(_safe_url(raw), "", "chuoi tuong doi phai bi loai (truoc fix)")
        fixed = self._f()([raw], self.HREF)[0]
        self.assertEqual(fixed, _safe_url(fixed))
        self.assertTrue(fixed.startswith("https://"))


class TestDocxFormatting(unittest.TestCase):
    def _one_question(self):
        return [{
            "exam_code": "TEST-1",
            "topic": 1,
            "question_num": 1,
            "question": "What is x?",
            "options": [
                {"letter": "A", "text": "alpha", "images": [], "is_correct": True},
                {"letter": "B", "text": "beta", "images": [], "is_correct": False},
            ],
            "suggested_answers": ["A"],
            "question_images": [],
        }]

    def _full_text(self, doc):
        parts = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        for s in doc.sections:
            for p in list(s.header.paragraphs) + list(s.footer.paragraphs):
                parts.append(p.text)
        return "\n".join(parts)

    def test_hide_answers_omits_answer_line(self):
        doc = Document()
        _add_question_docx(doc, self._one_question()[0], 1, hide_answers=True)
        self.assertNotIn("Đáp án:", self._full_text(doc))

    def test_shows_answer_line_by_default(self):
        doc = Document()
        _add_question_docx(doc, self._one_question()[0], 1)
        self.assertIn("Đáp án: A", self._full_text(doc))

    def test_default_font_is_set(self):
        doc = Document()
        _set_default_font(doc)
        # Ap dung cho ca style Heading, khong chi Normal.
        for style_name in ("Normal", "Heading 1", "Heading 2"):
            self.assertEqual(
                doc.styles[style_name].font.name, "Calibri",
                f"style {style_name} phai dung font Calibri",
            )

    def test_convert_adds_header_and_footer(self):
        import tempfile
        tmp = tempfile.mkdtemp()
        path = os.path.join(tmp, "t_questions.json")
        with open(path, "w", encoding="utf-8") as f:
            json.dump(self._one_question(), f)
        out = convert_to_docx(path, hide_answers=True)
        doc = Document(out)
        section = doc.sections[0]
        header_text = "\n".join(p.text for p in section.header.paragraphs)
        footer_text = "\n".join(p.text for p in section.footer.paragraphs)
        self.assertIn("TEST-1", header_text)
        self.assertIn("Trang", footer_text)
        shutil.rmtree(tmp)

    def test_answer_key_table_keeps_answers_when_hidden(self):
        doc = Document()
        _add_answer_key_table(doc, self._one_question())
        table = doc.tables[0]
        cell_texts = [c.text for row in table.rows for c in row.cells]
        self.assertIn("A", cell_texts)


class TestLoggingLevels(unittest.TestCase):
    """A1: logging phai giu nguyen output mac dinh, chi LOC khi --quiet va
    hien them khi --verbose. Day la diem Jev danh dau rui ro nhat, nen test
    kiem tra hanh vi quan sat duoc tren stream that.
    """

    def setUp(self):
        import examtopic.config as C
        self.C = C
        self.saved_level = C.LOG.level

    def tearDown(self):
        self.C.configure_logging()

    def _capture(self, **kwargs):
        stream = io.StringIO()
        self.C.configure_logging(stream=stream, **kwargs)
        self.C.LOG.info("INFO-marker")
        self.C.LOG.warning("WARNING-marker")
        self.C.LOG.debug("DEBUG-marker")
        return stream.getvalue()

    def test_default_shows_info_not_debug(self):
        out = self._capture()
        self.assertIn("INFO-marker", out)
        self.assertIn("WARNING-marker", out)
        self.assertNotIn("DEBUG-marker", out)

    def test_quiet_hides_info_keeps_warning(self):
        out = self._capture(quiet=True)
        self.assertNotIn("INFO-marker", out)
        self.assertIn("WARNING-marker", out)
        self.assertNotIn("DEBUG-marker", out)

    def test_verbose_shows_debug(self):
        out = self._capture(verbose=True)
        self.assertIn("INFO-marker", out)
        self.assertIn("DEBUG-marker", out)

    def test_output_has_no_logging_decorations(self):
        """Output phai la chuoi goc, KHONG co timestamp/level nhu logging mac dinh."""
        out = self._capture()
        self.assertNotIn("examtopic", out)
        self.assertNotIn("INFO:", out)
        self.assertNotIn("WARNING:", out)

    def test_message_text_is_byte_identical_to_old_print(self):
        """Bat bien quan trong nhat (Jev uoc tinh 0.61 kha nang pha vo): chuoi
        ma nguoi dung nhin thay phai giong HET print() cu, chi khac newline cuoi
        ma print() tu them. Neu ai do doi formatter lam hong dieu nay, test se bat.
        """
        message = "  Dang tai 3 anh (song song)..."
        stream = io.StringIO()
        self.C.configure_logging(stream=stream)
        self.C.LOG.info(message)
        # print(message) xuat ra dung `message + "\n"`.
        self.assertEqual(stream.getvalue(), message + "\n")

    def test_warning_and_error_have_no_prefix(self):
        """Canh bao/loi cung khong duoc tu them tien to nhu 'WARNING:'."""
        stream = io.StringIO()
        self.C.configure_logging(stream=stream)
        self.C.LOG.warning("  canh bao")
        self.C.LOG.error("  loi")
        self.assertEqual(stream.getvalue(), "  canh bao\n  loi\n")


class TestEnvConfig(unittest.TestCase):
    """A2: cac hang so phai doc duoc tu bien moi truong, va roi ve mac dinh
    an toan khi gia tri thieu/rac. Chay trong tien trinh con vi config duoc
    doc luc import.
    """

    def _read_config(self, env):
        import subprocess
        import sys as _sys
        code = ("import examtopic.config as C;"
                "print(C.MIN_DELAY, C.MAX_DELAY, C.RETRY_LIMIT, C.OUTPUT_DIR)")
        proc = subprocess.run(
            [_sys.executable, "-c", code],
            cwd=os.path.dirname(os.path.abspath(__file__)),
            capture_output=True, text=True, env={**os.environ, **env},
        )
        self.assertEqual(proc.returncode, 0, proc.stderr)
        return proc.stdout.split()

    def test_defaults_when_no_env(self):
        env = {k: v for k, v in os.environ.items()
               if k not in ("MIN_DELAY", "MAX_DELAY", "RETRY_LIMIT", "OUTPUT_DIR")}
        import subprocess, sys as _sys
        code = ("import examtopic.config as C;"
                "print(C.MIN_DELAY, C.MAX_DELAY, C.RETRY_LIMIT, C.OUTPUT_DIR)")
        proc = subprocess.run([_sys.executable, "-c", code],
                              cwd=os.path.dirname(os.path.abspath(__file__)),
                              capture_output=True, text=True, env=env)
        self.assertEqual(proc.stdout.split(), ["2", "5", "3", "output"])

    def test_env_overrides_applied(self):
        out = self._read_config({"MIN_DELAY": "7", "MAX_DELAY": "9",
                                 "RETRY_LIMIT": "5", "OUTPUT_DIR": "/tmp/zz"})
        self.assertEqual(out, ["7", "9", "5", "/tmp/zz"])

    def test_garbage_env_falls_back_to_default(self):
        out = self._read_config({"MIN_DELAY": "abc", "RETRY_LIMIT": "-3"})
        self.assertEqual(out[0], "2", "MIN_DELAY rac phai ve mac dinh")
        self.assertEqual(out[2], "3", "RETRY_LIMIT am phai ve mac dinh")


class TestSafeUrl(unittest.TestCase):
    """C2: chan scheme nguy hiem (javascript:, data:) truoc khi render href/src
    trong HTML xuat ra, de convert file JSON khong dang tin khong gay stored XSS.
    """

    def setUp(self):
        from examtopic.exporters.html import _safe_url
        self.safe = _safe_url

    def test_allows_http_and_https(self):
        self.assertEqual(self.safe("https://www.examtopics.com/x"), "https://www.examtopics.com/x")
        self.assertEqual(self.safe("http://a.com/i.png"), "http://a.com/i.png")
        self.assertEqual(self.safe("HTTPS://OK.COM/x"), "HTTPS://OK.COM/x")

    def test_blocks_dangerous_schemes(self):
        for bad in ("javascript:alert(1)", "data:text/html,<script>alert(1)</script>",
                    "vbscript:msgbox(1)", "file:///etc/passwd"):
            self.assertEqual(self.safe(bad), "", f"phai chan: {bad}")

    def test_rejects_empty_and_relative(self):
        for bad in ("", None, "   ", "/relative/path", "no-scheme.com/x"):
            self.assertEqual(self.safe(bad), "", f"phai loai: {bad!r}")

    def test_build_html_drops_javascript_href(self):
        from examtopic.exporters.html import build_html
        qs = [{"exam_code": "x", "topic": 1, "question_num": 1, "question": "q",
               "options": [], "suggested_answers": [], "answers": [],
               "url": "javascript:alert(document.cookie)"}]
        out = build_html(qs, "x", embed_images=False)
        self.assertNotIn('href="javascript:', out)
        self.assertNotIn('class="source-link"', out)

    def test_build_html_keeps_legit_href(self):
        from examtopic.exporters.html import build_html
        url = "https://www.examtopics.com/discussions/x/"
        qs = [{"exam_code": "x", "topic": 1, "question_num": 1, "question": "q",
               "options": [], "suggested_answers": [], "answers": [], "url": url}]
        out = build_html(qs, "x", embed_images=False)
        self.assertIn(f'href="{url}"', out)


class TestRecordTopicSafety(unittest.TestCase):
    """`topic` khong phai int (dict/list/None/chuoi) tung lam crash converter voi
    `TypeError: unhashable type: 'dict'` khi dua vao set(). Day la kieu du lieu
    co that khi convert file JSON chinh sua tay hoac tu nguon khac.
    """

    def test_record_topic_normalizes_all_types(self):
        from examtopic.parser import record_topic
        self.assertEqual(record_topic({"topic": 3}), 3)
        self.assertEqual(record_topic({"topic": "5"}), 5)
        self.assertEqual(record_topic({"topic": {"a": 1}}), 1)
        self.assertEqual(record_topic({"topic": [1, 2]}), 1)
        self.assertEqual(record_topic({"topic": None}), 1)
        self.assertEqual(record_topic({}), 1)
        self.assertEqual(record_topic("not a dict"), 1)

    def test_record_topic_result_is_hashable(self):
        from examtopic.parser import record_topic
        for bad in ({"a": 1}, [1, 2], None, "3", 3.7):
            set([record_topic({"topic": bad})])  # khong duoc raise

    def test_build_html_survives_unhashable_topic(self):
        from examtopic.exporters.html import build_html
        qs = [{"exam_code": "x", "topic": {"a": 1}, "question_num": 1, "question": "q",
               "options": [], "suggested_answers": [], "answers": []}]
        build_html(qs, "x", embed_images=False)  # truoc day crash tai day

    def test_docx_answer_key_survives_unhashable_topic(self):
        from docx import Document
        from examtopic.exporters.docx import _add_answer_key_table
        doc = Document()
        _add_answer_key_table(doc, [{"topic": {"a": 1}, "question_num": 1, "question": "q",
                                     "suggested_answers": ["A"]}])
        self.assertEqual(len(doc.tables), 1)

    def test_convert_survives_unhashable_topic(self):
        tmp = tempfile.mkdtemp()
        try:
            path = os.path.join(tmp, "x_questions.json")
            with open(path, "w", encoding="utf-8") as f:
                json.dump([{"exam_code": "X", "topic": [1], "question_num": 1,
                            "question": "q", "options": [], "suggested_answers": ["A"],
                            "answers": []}], f)
            convert_to_html(path)
            convert_to_docx(path)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


class TestQuietVerboseCli(unittest.TestCase):
    """A1: --quiet phai im lang voi convert thanh cong, nhung van in loi va
    tra ma loi khi that bai.
    """

    def _run(self, args):
        import subprocess
        import sys as _sys
        return subprocess.run(
            [_sys.executable, "tool.py", *args],
            cwd=os.path.dirname(os.path.abspath(__file__)),
            capture_output=True, text=True,
        )

    def _write_valid(self, tmp):
        path = os.path.join(tmp, "x_questions.json")
        with open(path, "w", encoding="utf-8") as f:
            json.dump([{"exam_code": "X", "topic": 1, "question_num": 1, "question": "q",
                        "options": [], "suggested_answers": [], "answers": []}], f)
        return path

    def test_quiet_success_is_silent(self):
        tmp = tempfile.mkdtemp()
        try:
            proc = self._run(["--quiet", "--convert-only", self._write_valid(tmp)])
            self.assertEqual(proc.returncode, 0, proc.stderr)
            self.assertEqual(proc.stdout.strip(), "", "chap nhan --quiet phai khong in gi")
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_quiet_still_reports_errors(self):
        proc = self._run(["--quiet", "--convert-only", "/nonexistent/x.json"])
        self.assertNotEqual(proc.returncode, 0)
        self.assertIn("Loi", proc.stdout + proc.stderr)

    def test_default_prints_conversion_paths(self):
        tmp = tempfile.mkdtemp()
        try:
            proc = self._run(["--convert-only", self._write_valid(tmp)])
            self.assertEqual(proc.returncode, 0, proc.stderr)
            self.assertIn("HTML:", proc.stdout)
            self.assertIn("DOCX:", proc.stdout)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":

    unittest.main()
